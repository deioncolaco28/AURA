"""
app/content/adapters/docx_adapter.py

DOCX document extraction, search, and generation adapter for AURA using python-docx.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import docx

from app.content.adapters.base import BaseContentAdapter
from app.content.models import (
    BlockType,
    ContentBlock,
    ContentDocument,
    ContentSection,
    ContentType,
)


class DOCXAdapter(BaseContentAdapter):
    """Extracts, indexes, searches, and generates Microsoft Word (.docx) documents."""

    def extract(self, source_path_or_url: str) -> ContentDocument:
        """Extract paragraphs, headings, and tables from a DOCX file."""
        path = Path(source_path_or_url).resolve()
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: '{path}'")

        try:
            doc = docx.Document(str(path))
        except Exception as e:
            raise RuntimeError(f"Failed to open DOCX '{path}': {e}")

        sections: list[ContentSection] = []
        current_section = ContentSection(title="Document Body", location="Section: 1")
        section_idx = 1

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else ""
            if "heading" in style_name:
                if current_section.blocks or current_section.raw_text:
                    sections.append(current_section)
                    section_idx += 1
                current_section = ContentSection(
                    title=text,
                    location=f"Heading: {text}",
                    blocks=[ContentBlock(block_type=BlockType.HEADING, text=text)],
                )
            else:
                current_section.blocks.append(
                    ContentBlock(block_type=BlockType.PARAGRAPH, text=text)
                )

        # Extract tables
        for table_idx, table in enumerate(doc.tables, start=1):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            table_text = "\n".join(table_rows)
            current_section.blocks.append(
                ContentBlock(
                    block_type=BlockType.TABLE,
                    text=table_text,
                    metadata={"table_index": table_idx},
                )
            )

        if current_section.blocks:
            sections.append(current_section)

        return ContentDocument(
            title=path.stem,
            source_type=ContentType.DOCX,
            source_path=str(path),
            sections=sections,
            metadata={"paragraphs_count": len(doc.paragraphs), "tables_count": len(doc.tables)},
        )

    def create_document(
        self,
        output_path: str | Path,
        title: str = "Document",
        content: str | list[str] | list[dict[str, Any]] = "",
        key_points: list[str] | None = None,
    ) -> str:
        """Generate a new formatted DOCX file."""
        target = Path(output_path).resolve()
        target.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document()
        doc.add_heading(title, level=0)

        if key_points:
            doc.add_heading("Key Highlights", level=1)
            for point in key_points:
                doc.add_paragraph(point, style="List Bullet")

        if isinstance(content, str) and content:
            for para in content.split("\n\n"):
                p_clean = para.strip()
                if p_clean:
                    doc.add_paragraph(p_clean)
        elif isinstance(content, list):
            for item in content:
                if isinstance(item, str):
                    doc.add_paragraph(item)
                elif isinstance(item, dict):
                    h = item.get("heading")
                    if h:
                        doc.add_heading(h, level=1)
                    body = item.get("body") or item.get("text", "")
                    if body:
                        doc.add_paragraph(body)

        doc.save(str(target))
        return str(target)
