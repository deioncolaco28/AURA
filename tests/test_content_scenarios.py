"""
Integration tests for Content Intelligence Scenarios and natural-language commands.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock
import docx
import openpyxl
import pytest

from app.automation.action_executor import ActionExecutor
from app.config.constants import AssistantState
from app.content.content_manager import ContentManager
from app.core.agent import Agent
from app.intelligence.action import Action, ActionType
from app.intelligence.planner import Planner
from app.intelligence.task_decomposer import RuleBasedTaskDecomposer


def _create_sample_doc(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Climate Report", level=0)
    doc.add_paragraph("Global temperatures have seen anomalous rises in urban micro-climates.")
    doc.save(str(path))


def _create_sample_sheet(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["City", "AQI"])
    ws.append(["Mumbai", 185])
    ws.append(["Delhi", 310])
    wb.save(str(path))


class TestContentScenarios:
    def test_scenario_1_summarize_document_command(self, tmp_path):
        # Scenario 1: "Summarize report.docx"
        doc_path = tmp_path / "report.docx"
        _create_sample_doc(doc_path)

        decomposer = RuleBasedTaskDecomposer()
        actions = decomposer.decompose(f"summarize {doc_path}")

        assert len(actions) == 2
        assert actions[0].action_type == ActionType.EXTRACT_DOCUMENT
        assert actions[0].target == str(doc_path)
        assert actions[1].action_type == ActionType.SUMMARIZE_DOCUMENT

    def test_scenario_2_summarize_and_save_docx(self, tmp_path):
        # Scenario 2: "summarize report.docx and save it as a word document"
        doc_path = tmp_path / "report.docx"
        _create_sample_doc(doc_path)

        decomposer = RuleBasedTaskDecomposer()
        actions = decomposer.decompose(f"summarize {doc_path} and save it as a word document")

        assert len(actions) == 3
        assert actions[0].action_type == ActionType.EXTRACT_DOCUMENT
        assert actions[1].action_type == ActionType.SUMMARIZE_DOCUMENT
        assert actions[2].action_type == ActionType.CREATE_DOCX

    def test_scenario_3_find_highest_value_in_spreadsheet(self, tmp_path):
        # Scenario 3: "find the highest AQI value in this spreadsheet"
        sheet_path = tmp_path / "data.xlsx"
        _create_sample_sheet(sheet_path)

        decomposer = RuleBasedTaskDecomposer()
        actions = decomposer.decompose(f"find the highest AQI value in {sheet_path}")

        assert len(actions) == 1
        assert actions[0].action_type == ActionType.ANALYZE_SHEET
        assert actions[0].target == str(sheet_path)

    def test_scenario_4_read_webpage_and_summarize(self):
        # Scenario 4: "read this webpage and summarize it"
        decomposer = RuleBasedTaskDecomposer()
        actions = decomposer.decompose("read this webpage and summarize it")

        assert len(actions) == 2
        assert actions[0].action_type == ActionType.EXTRACT_PAGE_CONTENT
        assert actions[1].action_type == ActionType.SUMMARIZE_DOCUMENT

    def test_scenario_5_turn_report_into_presentation(self, tmp_path):
        # Scenario 5: "turn this report into a presentation"
        decomposer = RuleBasedTaskDecomposer()
        actions = decomposer.decompose("turn this report into a presentation")

        assert len(actions) == 3
        assert actions[0].action_type == ActionType.EXTRACT_DOCUMENT
        assert actions[1].action_type == ActionType.SUMMARIZE_DOCUMENT
        assert actions[2].action_type == ActionType.CREATE_PPTX
