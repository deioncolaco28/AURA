"""
Tests for DOCXAdapter: extraction, headings, tables, search, and creation.
"""

from __future__ import annotations

from pathlib import Path
import docx
import pytest

from app.content.adapters.docx_adapter import DOCXAdapter
from app.content.models import BlockType, ContentType


def _create_sample_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Research Report", level=0)
    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("This study investigates closed-loop agent decision making.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "98.5%"
    
    doc.save(str(path))


class TestDOCXAdapter:
    def test_extract_docx_headings_and_tables(self, tmp_path):
        docx_path = tmp_path / "report.docx"
        _create_sample_docx(docx_path)

        adapter = DOCXAdapter()
        doc = adapter.extract(str(docx_path))

        assert doc.source_type == ContentType.DOCX
        assert doc.title == "report"
        assert len(doc.sections) >= 1
        
        full_text = doc.full_text
        assert "closed-loop agent" in full_text
        assert "Accuracy | 98.5%" in full_text

    def test_search_docx_with_section_provenance(self, tmp_path):
        docx_path = tmp_path / "report.docx"
        _create_sample_docx(docx_path)

        adapter = DOCXAdapter()
        doc = adapter.extract(str(docx_path))

        results = adapter.search(doc, "Methodology")
        assert len(results) >= 1
        assert "Heading: Methodology" in results[0].location or "report" in results[0].document_path

    def test_create_document(self, tmp_path):
        out_path = tmp_path / "output.docx"
        adapter = DOCXAdapter()
        
        created = adapter.create_document(
            output_path=out_path,
            title="Generated Summary",
            content="This is the main generated text.",
            key_points=["Point 1", "Point 2"],
        )

        assert Path(created).exists()
        doc_read = adapter.extract(created)
        assert "Generated Summary" in doc_read.full_text
        assert "Point 1" in doc_read.full_text
