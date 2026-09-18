"""
Tests for ContentManager, ContentSummarizer, and ContentQAEngine.
"""

from __future__ import annotations

from pathlib import Path
import docx
import pytest

from app.content.content_manager import ContentManager
from app.content.models import ContentType


def _create_test_report(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Environmental Impact Assessment", level=0)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This report evaluates atmospheric particulate pollution in urban centers.")
    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("We deployed laser sensor networks across ten major monitoring stations.")
    doc.add_paragraph("Data was collected continuously over a six month period.")
    doc.add_heading("Results", level=1)
    doc.add_paragraph("The average PM2.5 concentration was 58.2 micrograms per cubic meter.")
    doc.save(str(path))


class TestContentManager:
    def test_detect_type(self):
        cm = ContentManager()
        assert cm.detect_type("report.pdf") == ContentType.PDF
        assert cm.detect_type("report.docx") == ContentType.DOCX
        assert cm.detect_type("data.xlsx") == ContentType.XLSX
        assert cm.detect_type("slides.pptx") == ContentType.PPTX
        assert cm.detect_type("https://example.com/docs") == ContentType.WEB

    def test_summarize_document(self, tmp_path):
        report_path = tmp_path / "report.docx"
        _create_test_report(report_path)

        cm = ContentManager()
        summary = cm.summarize(str(report_path), max_points=3)

        assert len(summary.key_points) >= 1
        assert summary.word_count > 0
        assert summary.source_path == str(report_path)

    def test_grounded_qa_success_and_provenance(self, tmp_path):
        report_path = tmp_path / "report.docx"
        _create_test_report(report_path)

        cm = ContentManager()
        qa_res = cm.answer_question(str(report_path), "What methodology does this study use?")

        assert qa_res.is_grounded is True
        assert "laser sensor networks" in qa_res.answer.lower()
        assert "Source:" in qa_res.answer

    def test_grounded_qa_missing_information_fallback(self, tmp_path):
        report_path = tmp_path / "report.docx"
        _create_test_report(report_path)

        cm = ContentManager()
        qa_res = cm.answer_question(str(report_path), "What is the quantum computing algorithm?")

        assert qa_res.is_grounded is False
        assert "couldn't find enough information" in qa_res.answer.lower()

    def test_task_scoped_caching(self, tmp_path):
        report_path = tmp_path / "report.docx"
        _create_test_report(report_path)

        cm = ContentManager()
        doc1 = cm.extract(str(report_path))
        doc2 = cm.extract(str(report_path))

        assert doc1 is doc2  # Same cached instance

        cm.clear_cache()
        doc3 = cm.extract(str(report_path))
        assert doc3 is not doc1
