"""
Integration tests for Cross-Content Workflows and TaskGraph execution.
"""

from __future__ import annotations

from pathlib import Path
import docx
import openpyxl
import pytest

from app.automation.action_executor import ActionExecutor
from app.content.content_manager import ContentManager
from app.core.execution_engine import ExecutionEngine
from app.intelligence.action import Action, ActionType
from app.intelligence.expected_state import ExpectedState
from app.intelligence.task_graph import TaskGraph, TaskNode


def _create_source_report(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Autonomous Systems Report", level=0)
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("Autonomous assistants improve developer productivity by managing multi-step workflows.")
    doc.add_heading("Safety", level=1)
    doc.add_paragraph("Deterministic task graphs prevent unexpected or runaway actions.")
    doc.save(str(path))


def _create_source_sheet(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pollution"
    ws.append(["Station", "AQI"])
    ws.append(["Central", 180])
    ws.append(["North", 95])
    ws.append(["South", 220])
    wb.save(str(path))


class TestCrossContentWorkflows:
    def test_workflow_document_summary_to_docx(self, tmp_path):
        # Workflow: Extract document -> Summarize -> Create DOCX
        src_path = tmp_path / "source.docx"
        out_docx = tmp_path / "summary.docx"
        _create_source_report(src_path)

        graph = TaskGraph(goal="Summarize source and save as Word document")
        t1 = TaskNode(
            task_id="extract",
            action=Action(action_type=ActionType.EXTRACT_DOCUMENT, target=str(src_path)),
        )
        t2 = TaskNode(
            task_id="summarize",
            action=Action(action_type=ActionType.SUMMARIZE_DOCUMENT, target=str(src_path)),
            dependencies=["extract"],
        )
        t3 = TaskNode(
            task_id="create_docx",
            action=Action(
                action_type=ActionType.CREATE_DOCX,
                target=str(out_docx),
                value="$summary_text",
            ),
            dependencies=["summarize"],
            expected_state=ExpectedState(file_exists=str(out_docx)),
        )

        for n in (t1, t2, t3):
            graph.add_node(n)

        cm = ContentManager()
        executor = ActionExecutor(content_manager=cm)

        engine = ExecutionEngine(
            execute_action=executor.execute,
            verify_action=lambda a: None,
        )

        ctx = engine.run_graph(graph)
        assert graph.is_completed() is True
        assert out_docx.exists()

        created_doc = cm.extract(str(out_docx))
        assert "Autonomous assistants" in created_doc.full_text or "productivity" in created_doc.full_text

    def test_workflow_document_summary_to_pptx(self, tmp_path):
        # Workflow: Extract document -> Summarize -> Create Presentation
        src_path = tmp_path / "source.docx"
        out_pptx = tmp_path / "deck.pptx"
        _create_source_report(src_path)

        graph = TaskGraph(goal="Create presentation from report")
        t1 = TaskNode(
            task_id="extract",
            action=Action(action_type=ActionType.EXTRACT_DOCUMENT, target=str(src_path)),
        )
        t2 = TaskNode(
            task_id="summarize",
            action=Action(action_type=ActionType.SUMMARIZE_DOCUMENT, target=str(src_path)),
            dependencies=["extract"],
        )
        t3 = TaskNode(
            task_id="create_pptx",
            action=Action(
                action_type=ActionType.CREATE_PPTX,
                target=str(out_pptx),
                parameters={"title": "Autonomous Systems", "key_points": ["Point A", "Point B"]},
            ),
            dependencies=["summarize"],
            expected_state=ExpectedState(file_exists=str(out_pptx)),
        )

        for n in (t1, t2, t3):
            graph.add_node(n)

        cm = ContentManager()
        executor = ActionExecutor(content_manager=cm)

        engine = ExecutionEngine(
            execute_action=executor.execute,
            verify_action=lambda a: None,
        )

        ctx = engine.run_graph(graph)
        assert graph.is_completed() is True
        assert out_pptx.exists()

    def test_workflow_xlsx_analysis_and_export(self, tmp_path):
        # Workflow: Analyze spreadsheet -> Extract high values -> Create new XLSX
        sheet_path = tmp_path / "stations.xlsx"
        out_sheet = tmp_path / "filtered.xlsx"
        _create_source_sheet(sheet_path)

        cm = ContentManager()
        stats = cm.analyze_spreadsheet(str(sheet_path))
        assert stats["numeric_columns"]["AQI"]["max"] == 220

        created = cm.create_document(
            "XLSX",
            str(out_sheet),
            headers=["Station", "High_AQI"],
            rows=[["South", 220]],
        )
        assert Path(created).exists()
